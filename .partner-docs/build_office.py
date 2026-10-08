#!/usr/bin/env python3
"""Build editable DOCX + PPTX from the partnership Markdown sources.

Parses the same Markdown the PDF is built from, so the three formats can never
drift. Arabic runs are marked RTL with the project's bundled Noto Sans Arabic.
"""
import pathlib
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.util import Pt as PPt, Inches as PIn

ROOT = pathlib.Path('/workspace/project')
DOCS = ROOT / '.partner-docs'
FONT = 'Noto Sans Arabic'

ARABIC = re.compile(r'[\u0600-\u06FF]')
NAVY = RGBColor(0x0F, 0x3D, 0x5C)
GREEN = RGBColor(0x14, 0x53, 0x2D)
GREY = RGBColor(0x33, 0x41, 0x55)
PNAVY = PRGB(0x0F, 0x3D, 0x5C)
PGREEN = PRGB(0x14, 0x53, 0x2D)
PGREY = PRGB(0x33, 0x41, 0x55)
PWHITE = PRGB(0xFF, 0xFF, 0xFF)


def rtl(par):
    p = par._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    p.append(bidi)
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def run(par, text, size=10.5, bold=False, color=GREY, font=FONT):
    r = par.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    rf.set(qn('w:cs'), font)
    rf.set(qn('w:ascii'), font)
    rf.set(qn('w:hAnsi'), font)
    if ARABIC.search(text):
        rtl_el = OxmlElement('w:rtl')
        rpr.append(rtl_el)
        cs = OxmlElement('w:cs')
        rpr.append(cs)
    return r


def parse_blocks(text):
    """Yield (kind, payload) for the Markdown constructs we render."""
    lines = text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('```'):
            buf = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            yield 'code', '\n'.join(buf)
            continue
        m = re.match(r'^(#{1,4})\s+(.*)$', line)
        if m:
            yield 'h' + str(len(m.group(1))), m.group(2).strip()
            i += 1
            continue
        if line.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:|-]+\|$', lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                rows.append(cells)
                i += 1
            yield 'table', rows
            continue
        if line.strip() in ('---', '***'):
            yield 'hr', ''
            i += 1
            continue
        if line.startswith('> '):
            buf = []
            while i < len(lines) and lines[i].startswith('>'):
                buf.append(lines[i].lstrip('> ').strip())
                i += 1
            yield 'quote', ' '.join(buf)
            continue
        if re.match(r'^[-*]\s+', line):
            buf = []
            while i < len(lines) and re.match(r'^[-*]\s+', lines[i]):
                buf.append(re.sub(r'^[-*]\s+', '', lines[i]).strip())
                i += 1
            yield 'ul', buf
            continue
        if line.strip() == '':
            i += 1
            continue
        yield 'p', line.strip()
        i += 1


def strip_md(s):
    s = re.sub(r'\*\*(.+?)\*\*', r'\1', s)
    s = re.sub(r'`(.+?)`', r'\1', s)
    s = re.sub(r'\[(.+?)\]\(.+?\)', r'\1', s)
    return s


# ---------------- DOCX ----------------
def build_docx(src, out, cover_lines, title):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.75)
    sec.left_margin = sec.right_margin = Inches(0.7)

    def page():
        doc.add_page_break()

    # Cover
    for text, size, color, bold in cover_lines:
        p = doc.add_paragraph()
        rtl(p)
        run(p, text, size=size, bold=bold, color=color)
    page()

    for kind, payload in parse_blocks((DOCS / src).read_text(encoding='utf-8')):
        if kind in ('h1', 'h2', 'h3', 'h4'):
            lvl = int(kind[1])
            sizes = {1: 19, 2: 14.5, 3: 12.5, 4: 11}
            colors = {1: NAVY, 2: NAVY, 3: GREEN, 4: GREY}
            p = doc.add_paragraph()
            rtl(p)
            if lvl == 2:
                page()
            run(p, strip_md(payload), size=sizes[lvl], bold=True, color=colors[lvl])
        elif kind == 'p':
            p = doc.add_paragraph()
            rtl(p)
            run(p, strip_md(payload))
        elif kind == 'ul':
            for item in payload:
                p = doc.add_paragraph(style='List Bullet')
                rtl(p)
                run(p, strip_md(item))
        elif kind == 'quote':
            p = doc.add_paragraph()
            rtl(p)
            run(p, strip_md(payload), size=9.8, color=RGBColor(0x92, 0x40, 0x0E))
        elif kind == 'code':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(payload)
            r.font.name = 'DejaVu Sans Mono'
            r.font.size = Pt(8)
        elif kind == 'table':
            rows = payload
            t = doc.add_table(rows=0, cols=len(rows[0]))
            t.style = 'Light Grid Accent 1'
            t.alignment = WD_TABLE_ALIGNMENT.RIGHT
            for ri, cells in enumerate(rows):
                cell_row = t.add_row().cells
                for ci, cell in enumerate(cells):
                    if ci >= len(cell_row):
                        break
                    par = cell_row[ci].paragraphs[0]
                    rtl(par)
                    if ri == 0:
                        run(par, strip_md(cell), size=9, bold=True,
                            color=RGBColor(0xFF, 0xFF, 0xFF))
                    else:
                        run(par, strip_md(cell), size=9)
        elif kind == 'hr':
            doc.add_paragraph('─' * 40)
    doc.save(DOCS / out)
    print('docx', out)


# ---------------- PPTX ----------------
def build_pptx(src, out, title):
    prs = Presentation()
    prs.slide_width = PIn(13.333)
    prs.slide_height = PIn(7.5)
    blank = prs.slide_layouts[6]

    def slide(title_text, bullets, subtitle=None):
        s = prs.slides.add_slide(blank)
        # Title
        tb = s.shapes.add_textbox(PIn(0.6), PIn(0.35), PIn(12.1), PIn(1.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = 3  # right
        r = p.add_run()
        r.text = title_text
        r.font.size = PPt(28)
        r.font.bold = True
        r.font.color.rgb = PNAVY
        r.font.name = FONT
        y = 1.5
        if subtitle:
            sb = s.shapes.add_textbox(PIn(0.6), PIn(1.35), PIn(12.1), PIn(0.5))
            sp = sb.text_frame.paragraphs[0]
            sp.alignment = 3
            sr = sp.add_run()
            sr.text = subtitle
            sr.font.size = PPt(14)
            sr.font.color.rgb = PGREEN
            sr.font.name = FONT
            y = 2.0
        body = s.shapes.add_textbox(PIn(0.6), PIn(y), PIn(12.1), PIn(7.5 - y - 0.4))
        bf = body.text_frame
        bf.word_wrap = True
        first = True
        for b in bullets:
            p = bf.paragraphs[0] if first else bf.add_paragraph()
            first = False
            p.alignment = 3
            p.space_after = PPt(8)
            r = p.add_run()
            r.text = '•  ' + strip_md(b) if b else ''
            r.font.size = PPt(15)
            r.font.color.rgb = PGREY
            r.font.name = FONT
        return s

    # Title slide
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(PIn(1), PIn(2.2), PIn(11.3), PIn(3))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = [
        ('WASIT BUSINESS GLOBAL', 40, PNAVY, True),
        ('الوسيط للتجارة الرقمية', 26, PGREEN, True),
        ('Global Digital Trade Intermediary', 16, PGREY, False),
        ('Partnership Discussion — نقاش الشراكة', 15, PGREY, False),
    ]
    for idx, (t, sz, col, bold) in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = 2
        r = p.add_run()
        r.text = t
        r.font.size = PPt(sz)
        r.font.bold = bold
        r.font.color.rgb = col
        r.font.name = FONT

    current_title = None
    current_sub = None
    bullets = []

    def flush():
        if current_title:
            slide(current_title, bullets, current_sub)

    for kind, payload in parse_blocks((DOCS / src).read_text(encoding='utf-8')):
        if kind in ('h1', 'h2'):
            flush()
            current_title = strip_md(payload)
            current_sub = None
            bullets = []
        elif kind == 'h3':
            current_sub = strip_md(payload)
        elif kind == 'p':
            bullets.append(payload)
        elif kind == 'ul':
            bullets.extend(payload)
        elif kind == 'quote':
            bullets.append(payload)
        elif kind == 'table':
            for row in payload[1:]:
                bullets.append(' — '.join(strip_md(c) for c in row if c))
    flush()
    prs.save(DOCS / out)
    print('pptx', out)


if __name__ == '__main__':
    build_docx(
        'WASIT-BUSINESS-GLOBAL-PARTNERSHIP-DOSSIER.md',
        'WASIT-BUSINESS-GLOBAL-PARTNERSHIP-DOSSIER.docx',
        [
            ('WASIT BUSINESS GLOBAL', 26, NAVY, True),
            ('PARTNERSHIP PRESENTATION & PLATFORM DOSSIER', 15, GREEN, True),
            ('الوسيط للتجارة الرقمية', 20, NAVY, True),
            ('Global Digital Trade Intermediary', 12, GREY, False),
            ('الإصدار الموثّق: b03d17b · 2026-10-04', 10, GREY, False),
        ],
        'WASIT BUSINESS GLOBAL — Partnership Dossier',
    )
    build_docx(
        'WASIT-BUSINESS-GLOBAL-PARTNER-PRESENTATION.md',
        'WASIT-BUSINESS-GLOBAL-PARTNER-PRESENTATION.docx',
        [
            ('WASIT BUSINESS GLOBAL', 26, NAVY, True),
            ('PARTNER PRESENTATION', 15, GREEN, True),
            ('الوسيط للتجارة الرقمية', 20, NAVY, True),
            ('Global Digital Trade Intermediary', 12, GREY, False),
            ('الإصدار الموثّق: b03d17b · 2026-10-04', 10, GREY, False),
        ],
        'WASIT BUSINESS GLOBAL — Partner Presentation',
    )
    build_pptx(
        'WASIT-BUSINESS-GLOBAL-PARTNER-PRESENTATION.md',
        'WASIT-BUSINESS-GLOBAL-PARTNER-PRESENTATION.pptx',
        'WASIT BUSINESS GLOBAL — Partner Presentation',
    )
